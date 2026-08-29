"""Unit tests for main.py pure-logic helpers.

Run:  python3 test_main.py -v
"""
import sys
import tempfile
import unittest
from pathlib import Path

import fitz

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document

import main as itr


def _make_pdf_bytes(page_texts, toc=None):
    doc = fitz.open()
    try:
        for text in page_texts:
            page = doc.new_page()
            page.insert_text((72, 72), text)
        if toc is not None:
            doc.set_toc(toc)
        return doc.write()
    finally:
        doc.close()


def _page_texts(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        return [page.get_text().strip() for page in doc]
    finally:
        doc.close()


def _jur(name, abbr=None, label=None, outcome="no_tax", balance_due=None,
         overpayment=None, quote="q", note=None):
    """Facts-shape jurisdiction fixture (mirrors tests/test_email_rendering.py)."""
    return {"jurisdiction_name": name, "state_abbreviation": abbr,
            "display_label": label, "outcome": outcome,
            "balance_due": balance_due, "overpayment": overpayment,
            "source_quote": quote, "other_note": note}


def _pte_entry(amount=10500, date="06/15/2026", ordinal=None, **over):
    """Facts-shape ``scheduled_payments`` PTE entry fixture."""
    entry = {"type": "pte", "jurisdiction": "California", "amount": amount,
             "date": date, "payment_method": "direct_debit", "ordinal": ordinal,
             "source_quote": "q", "note": None}
    entry.update(over)
    return entry


def _docx_data(jurisdictions, return_type="Individual (1040)", firm=None):
    return {
        "client": {"name": "C"},
        "cpa_firm": firm or {"name": "CNY LLP", "sharefile_subdomain": "cnyllp"},
        "tax_year": "2025",
        "return_type": return_type,
        "jurisdictions": list(jurisdictions),
        "scheduled_payments": [],
    }


def _docx_text(path):
    doc = Document(path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


class TestMarkdownParagraph(unittest.TestCase):
    def setUp(self):
        self.doc = Document()

    def test_plain_text_has_one_non_bold_run(self):
        para = itr._add_markdown_paragraph(self.doc, "Hello world.")
        self.assertEqual(len(para.runs), 1)
        self.assertEqual(para.runs[0].text, "Hello world.")
        self.assertFalse(para.runs[0].font.bold)

    def test_single_bold_segment_alternates(self):
        para = itr._add_markdown_paragraph(self.doc, "Refund of **$2,052** deposited.")
        texts = [r.text for r in para.runs]
        bolds = [bool(r.font.bold) for r in para.runs]
        self.assertEqual(texts, ["Refund of ", "$2,052", " deposited."])
        self.assertEqual(bolds, [False, True, False])

    def test_multiple_bold_segments(self):
        text = "**Overpayment** of **$3,597** of which **$3,597** credited to 2026 tax."
        para = itr._add_markdown_paragraph(self.doc, text)
        bolds = [bool(r.font.bold) for r in para.runs]
        texts = [r.text for r in para.runs]
        # Leading "" segment is dropped; runs are: Overpayment, " of ", $3,597, " of which ", $3,597, " credited to 2026 tax."
        self.assertEqual(texts[0], "Overpayment")
        self.assertTrue(bolds[0])
        self.assertIn("$3,597", texts)

    def test_no_literal_asterisks_in_rendered_runs(self):
        para = itr._add_markdown_paragraph(self.doc, "A **B** C")
        for run in para.runs:
            self.assertNotIn("**", run.text)


class TestExtractPage1Bytes(unittest.TestCase):
    SAMPLE = Path(__file__).parent / "samples" / "Yamane, Jon H & Gail ITR 2025.pdf"

    def test_returns_non_empty_bytes(self):
        data = itr.extract_page1_bytes(str(self.SAMPLE))
        self.assertIsInstance(data, bytes)
        self.assertGreater(len(data), 1000)

    def test_result_is_a_one_page_pdf(self):
        import fitz
        data = itr.extract_page1_bytes(str(self.SAMPLE))
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            self.assertEqual(doc.page_count, 1)
        finally:
            doc.close()

    def test_first_page_text_preserved(self):
        import fitz
        data = itr.extract_page1_bytes(str(self.SAMPLE))
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            text = doc[0].get_text()
        finally:
            doc.close()
        self.assertIn("CNY LLP", text)
        self.assertIn("YAMANE", text)


class TestExtractCoverLetterBytes(unittest.TestCase):
    def test_extracts_normalized_letter_section_until_next_later_top_level(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "letter-page-1", "letter-page-2", "tax-summary"],
            [
                [1, "Front Matter", 1],
                [1, "  lEtTeR  ", 2],
                [2, "Page 1", 2],
                [1, "Invoice", 2],
                [1, "Tax Summary", 4],
            ],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["letter-page-1", "letter-page-2"])

    def test_missing_letter_bookmark_falls_back_to_first_page(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "tax-summary"],
            [[1, "Tax Summary", 2]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["physical-page-1"])

    def test_substring_title_does_not_match_letter(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "engagement-letter"],
            [[1, "Engagement Letter", 2]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["physical-page-1"])

    def test_nested_letter_bookmark_does_not_match(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "nested-letter"],
            [[1, "General Info.", 1], [2, "Letter", 2]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["physical-page-1"])

    def test_invalid_letter_destination_falls_back_to_first_page(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "tax-summary"],
            [[1, "Letter", -1], [1, "Tax Summary", 2]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["physical-page-1"])

    def test_skips_invalid_letter_and_uses_next_valid_exact_letter(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "valid-letter", "tax-summary"],
            [[1, "Letter", -1], [1, "Letter", 2], [1, "Tax Summary", 3]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["valid-letter"])

    def test_final_letter_bookmark_extracts_only_its_start_page(self):
        source = _make_pdf_bytes(
            ["physical-page-1", "letter-start", "unbounded-return-page"],
            [[1, "General Info.", 1], [1, "Letter", 2], [2, "Page 1", 2]],
        )

        result = itr.extract_cover_letter_bytes(source)

        self.assertEqual(_page_texts(result), ["letter-start"])


class TestApplyFtbFirstPteOrdinal(unittest.TestCase):
    FORM_PAGE = """California e-file Return Authorization for Corporations
FORM 8453-C
Part IV Pass-Through Entity (PTE) Elective Tax Payment for Taxable Year 2026 (for Form 100S only)
First Payment
9 Amount 10,500.
10 Withdrawal date
(mm/dd/yyyy) 6/15/2026
Part V Banking Information
"""

    def _task2_data(self, amount=10500, date="06/15/2026", return_type="S-Corporation (1120S)"):
        return {
            "return_type": return_type,
            "scheduled_payments": [_pte_entry(amount=amount, date=date)],
        }

    def test_adds_first_only_when_form_amount_and_date_match(self):
        pdf_bytes = _make_pdf_bytes([self.FORM_PAGE])
        task2_data = self._task2_data()

        result = itr.apply_ftb_first_pte_ordinal(pdf_bytes, task2_data)

        self.assertEqual(result["scheduled_payments"][0]["ordinal"], "1st")
        self.assertIsNone(task2_data["scheduled_payments"][0]["ordinal"])

    def test_supports_filled_ftb_llc_first_payment_for_partnership_return(self):
        llc_form = (
            self.FORM_PAGE
            .replace(
                "California e-file Return Authorization for Corporations",
                "California e-file Return Authorization for Limited Liability Companies",
            )
            .replace("8453-C", "8453-LLC")
            .replace("9 Amount", "8 Amount")
            .replace("10 Withdrawal date", "9 Withdrawal date")
        )

        result = itr.apply_ftb_first_pte_ordinal(
            _make_pdf_bytes([llc_form]),
            self._task2_data(return_type="Partnership (1065)"),
        )

        self.assertEqual(result["scheduled_payments"][0]["ordinal"], "1st")

    def test_supports_official_ftb_8453_p_partnership_section(self):
        partnership_form = (
            self.FORM_PAGE
            .replace(
                "California e-file Return Authorization for Corporations",
                "California e-file Return Authorization for Partnerships",
            )
            .replace("8453-C", "8453-P")
            .replace("Part IV", "Part III")
            .replace("9 Amount", "6 Amount")
            .replace("10 Withdrawal date", "7 Withdrawal date")
            .replace("Part V", "Part IV")
        )

        result = itr.apply_ftb_first_pte_ordinal(
            _make_pdf_bytes([partnership_form]),
            self._task2_data(return_type="Partnership (1065)"),
        )

        self.assertEqual(result["scheduled_payments"][0]["ordinal"], "1st")

    def test_amount_or_date_mismatch_does_not_add_first(self):
        pdf_bytes = _make_pdf_bytes([self.FORM_PAGE])
        cases = (
            {"amount": 10501, "date": "06/15/2026"},
            {"amount": 10500, "date": "06/16/2026"},
        )

        for case in cases:
            with self.subTest(case=case):
                result = itr.apply_ftb_first_pte_ordinal(
                    pdf_bytes,
                    self._task2_data(**case),
                )
                self.assertIsNone(result["scheduled_payments"][0]["ordinal"])

    def test_blank_or_unofficial_first_payment_text_does_not_add_first(self):
        pages = (
            self.FORM_PAGE.replace("10,500.", "").replace("6/15/2026", ""),
            "First Payment PTE amount $10,500 date 6/15/2026",
        )

        for page_text in pages:
            with self.subTest(page_text=page_text):
                result = itr.apply_ftb_first_pte_ordinal(
                    _make_pdf_bytes([page_text]),
                    self._task2_data(),
                )
                self.assertIsNone(result["scheduled_payments"][0]["ordinal"])

    def test_existing_ordinal_and_ineligible_return_are_unchanged(self):
        pdf_bytes = _make_pdf_bytes([self.FORM_PAGE])
        already_first = {
            "return_type": "S-Corporation (1120S)",
            "scheduled_payments": [_pte_entry(ordinal="1st")],
        }
        individual = self._task2_data(return_type="Individual (1040)")

        existing_result = itr.apply_ftb_first_pte_ordinal(pdf_bytes, already_first)
        ineligible_result = itr.apply_ftb_first_pte_ordinal(pdf_bytes, individual)

        self.assertEqual(existing_result, already_first)
        self.assertEqual(ineligible_result, individual)

    def test_duplicate_matching_pte_rows_are_left_unchanged(self):
        pdf_bytes = _make_pdf_bytes([self.FORM_PAGE])
        task2_data = self._task2_data()
        task2_data["scheduled_payments"].append(dict(task2_data["scheduled_payments"][0]))

        result = itr.apply_ftb_first_pte_ordinal(pdf_bytes, task2_data)

        self.assertEqual(result, task2_data)
        self.assertTrue(
            all(payment["ordinal"] is None for payment in result["scheduled_payments"])
        )

    def test_only_matching_row_changes_when_other_pte_payment_differs(self):
        pdf_bytes = _make_pdf_bytes([self.FORM_PAGE])
        task2_data = self._task2_data()
        other_entry = _pte_entry(amount=4000, date="09/15/2026")
        task2_data["scheduled_payments"].append(other_entry)

        result = itr.apply_ftb_first_pte_ordinal(pdf_bytes, task2_data)

        self.assertEqual(result["scheduled_payments"][0]["ordinal"], "1st")
        self.assertIsNone(result["scheduled_payments"][1]["ordinal"])


class TestGenerateEmailDocx(unittest.TestCase):
    """DOCX-render regression tests (moved here from tests/test_email_rendering.py
    when Task 6 split the HTML/DOCX renderers; fixtures updated to the facts shape)."""

    def _im_data(self):
        return _docx_data([
            _jur("Federal", outcome="refund_or_credit", overpayment={"refunded": 8647}),
            _jur("Oregon", abbr="OR", outcome="refund_or_credit", overpayment={"refunded": 7447}),
            _jur("Oregon Metro Supportive Housing Services", abbr="OR",
                 label="OR Metro Income Tax", outcome="refund_or_credit",
                 overpayment={"refunded": 349}),
        ])

    def _hong_data(self):
        return _docx_data([
            _jur("Federal", outcome="refund_or_credit", overpayment={"refunded": 951}),
            _jur("California", abbr="CA", outcome="refund_or_credit", overpayment={"refunded": 297}),
        ])

    def test_docx_combines_cny_brand_and_correct_im_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = str(Path(tmp) / "im.docx")
            itr.generate_email_docx(self._im_data(), output)
            text = _docx_text(output)

        federal_line = (
            "Federal Income Tax: No tax is payable with the filing of this return. "
            "Refund of $8,647 will be deposited to your account."
        )
        state_line = (
            "State Income Tax: No tax is payable with the filing of this return. "
            "Refund of $7,447 will be deposited to your account."
        )
        metro_line = (
            "OR Metro Income Tax: No tax is payable with the filing of this return. "
            "Refund of $349 will be deposited to your account."
        )

        self.assertIn("Your 2025 Income Tax Return and CNY invoice are now available", text)
        self.assertIn(federal_line, text)
        self.assertIn(state_line, text)
        self.assertIn(metro_line, text)
        self.assertLess(text.index(federal_line), text.index(state_line))
        self.assertLess(text.index(state_line), text.index(metro_line))
        self.assertNotIn("OR State Income Tax", text)
        self.assertIn('Enter "cnyllp" as the subdomain', text)

    def test_hong_refunds_remain_associated_with_the_correct_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = str(Path(tmp) / "hong.docx")
            itr.generate_email_docx(self._hong_data(), output)
            text = _docx_text(output)

        self.assertIn(
            "Federal Income Tax: No tax is payable with the filing of this return. "
            "Refund of $951 will be deposited to your account.",
            text,
        )
        self.assertIn(
            "State Income Tax: No tax is payable with the filing of this return. "
            "Refund of $297 will be deposited to your account.",
            text,
        )

    def test_docx_renders_bold_review_paragraphs_when_sentence_functions_return_none(self):
        data = _docx_data([
            _jur("Federal", outcome="other", note="Client must pay by money order",
                 quote="Mail a money order to..."),
            _jur("California", abbr="CA", outcome="refund_or_credit",
                 overpayment={"refunded": 100}),
        ])
        data["scheduled_payments"] = [
            {"type": "other", "jurisdiction": "California", "amount": 1,
             "date": "04/15/2026", "payment_method": "unspecified",
             "ordinal": None, "source_quote": "q", "note": "unclear line item"},
        ]

        with tempfile.TemporaryDirectory() as tmp:
            output = str(Path(tmp) / "review.docx")
            itr.generate_email_docx(data, output)
            paragraphs = list(Document(output).paragraphs)

        federal_review = next(p for p in paragraphs if p.text.startswith("Federal Income Tax: "))
        other_review = next(p for p in paragraphs if "unclear line item" in p.text)

        self.assertIn("NEEDS REVIEW", federal_review.text)
        self.assertIn("Client must pay by money order", federal_review.text)
        self.assertIn('Mail a money order to..."', federal_review.text)
        self.assertTrue(federal_review.runs)
        self.assertTrue(all(run.font.bold for run in federal_review.runs))

        self.assertIn("NEEDS REVIEW", other_review.text)
        self.assertTrue(other_review.runs)
        self.assertTrue(all(run.font.bold for run in other_review.runs))

        # California's valid refund sentence must still render normally alongside the review block.
        state_text = "\n".join(p.text for p in paragraphs)
        self.assertIn(
            "State Income Tax: No tax is payable with the filing of this return. "
            "Refund of $100 will be deposited to your account.",
            state_text,
        )


if __name__ == "__main__":
    unittest.main()
