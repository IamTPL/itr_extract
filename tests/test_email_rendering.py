import copy

import pytest
from docx import Document

from jobs.email_html import generate_email_html
from main import generate_email_docx


def _base_data(firm_name="CNY LLP"):
    return {
        "cpa_firm": {"name": firm_name, "sharefile_subdomain": "cnyllp"},
        "tax_year": "2025",
        "return_type": "Individual (1040)",
        "tax_summary": {"federal_sentence": None, "state_sentences": []},
        "estimated_payments": [],
        "pte_payments": [],
    }


def _im_data():
    data = _base_data()
    data["tax_summary"] = {
        "federal_sentence": (
            "No tax is payable with the filing of this return. "
            "**Refund** of **$8,647** will be deposited to your account."
        ),
        "state_sentences": [
            {
                "state_name": "Oregon",
                "state_abbreviation": "OR",
                "display_label": None,
                "sentence": (
                    "No tax is payable with the filing of this return. "
                    "**Refund** of **$7,447** will be deposited to your account."
                ),
            },
            {
                "state_name": "Oregon Metro Supportive Housing Services",
                "state_abbreviation": "OR",
                "display_label": "OR Metro Income Tax",
                "sentence": (
                    "No tax is payable with the filing of this return. "
                    "**Refund** of **$349** will be deposited to your account."
                ),
            },
        ],
    }
    return data


def _hong_data():
    data = _base_data()
    data["tax_summary"] = {
        "federal_sentence": (
            "No tax is payable with the filing of this return. "
            "**Refund** of **$951** will be deposited to your account."
        ),
        "state_sentences": [
            {
                "state_name": "California",
                "state_abbreviation": "CA",
                "display_label": None,
                "sentence": (
                    "No tax is payable with the filing of this return. "
                    "**Refund** of **$297** will be deposited to your account."
                ),
            }
        ],
    }
    return data


def test_html_combines_cny_brand_and_correct_im_labels():
    html = generate_email_html(_im_data())

    federal_line = (
        "Federal Income Tax: No tax is payable with the filing of this return. "
        "<strong>Refund</strong> of <strong>$8,647</strong> will be deposited to your account."
    )
    state_line = (
        "State Income Tax: No tax is payable with the filing of this return. "
        "<strong>Refund</strong> of <strong>$7,447</strong> will be deposited to your account."
    )
    metro_line = (
        "OR Metro Income Tax: No tax is payable with the filing of this return. "
        "<strong>Refund</strong> of <strong>$349</strong> will be deposited to your account."
    )

    assert "Your 2025 Income Tax Return and CNY invoice are now available" in html
    assert federal_line in html
    assert state_line in html
    assert metro_line in html
    assert html.index(federal_line) < html.index(state_line) < html.index(metro_line)
    assert "OR State Income Tax" not in html
    assert "Income Tax Income Tax" not in html
    assert "OR Metro Income Tax::" not in html
    assert "Enter &ldquo;cnyllp&rdquo; as the subdomain" in html


@pytest.mark.parametrize("firm", ["CNY LLP", "ABC LLP", None, ""])
def test_html_invoice_brand_is_global_and_does_not_mutate_input(firm):
    data = _base_data(firm)
    before = copy.deepcopy(data)

    html = generate_email_html(data)

    assert "Your 2025 Income Tax Return and CNY invoice are now available" in html
    assert "CNY LLP invoice" not in html
    assert "ABC LLP invoice" not in html
    assert "incoive" not in html
    assert data == before


@pytest.mark.parametrize("firm_value", ["missing", None])
def test_html_handles_missing_firm_object(firm_value):
    data = _base_data()
    if firm_value == "missing":
        data.pop("cpa_firm")
    else:
        data["cpa_firm"] = None

    html = generate_email_html(data)

    assert "CNY invoice" in html
    assert "Enter the firm subdomain" in html


def test_html_handles_firm_object_without_name():
    data = _base_data()
    data["cpa_firm"].pop("name")

    html = generate_email_html(data)

    assert "CNY invoice" in html
    assert "Enter &ldquo;cnyllp&rdquo; as the subdomain" in html


def test_html_escapes_model_produced_display_label():
    data = _base_data()
    data["tax_summary"]["state_sentences"] = [
        {
            "state_name": "Local",
            "state_abbreviation": "OR",
            "display_label": "A&B <Local> Income Tax",
            "sentence": "**Refund** of **$1**.",
        }
    ]

    html = generate_email_html(data)

    assert "A&amp;B &lt;Local&gt; Income Tax:" in html
    assert "A&B <Local> Income Tax:" not in html


def test_html_escapes_all_dynamic_text_without_losing_markdown_bold():
    data = _base_data()
    data["tax_year"] = "<img src=x onerror=alert(1)>"
    data["cpa_firm"]["sharefile_subdomain"] = "x</li><script>alert(2)</script>"
    data["tax_summary"]["federal_sentence"] = (
        "<svg onload=alert(3)> **Refund** of **$1**."
    )
    data["estimated_payments"] = [
        {
            "date": "<a href=evil>04/15/2026</a>",
            "federal": 1,
            "state": 0,
            "state_name": None,
        }
    ]

    html = generate_email_html(data)

    assert "<img" not in html
    assert "<script" not in html
    assert "<svg" not in html
    assert "<a href=evil" not in html
    assert "&lt;img src=x onerror=alert(1)&gt;" in html
    assert "x&lt;/li&gt;&lt;script&gt;alert(2)&lt;/script&gt;" in html
    assert "&lt;svg onload=alert(3)&gt; <strong>Refund</strong>" in html
    assert "&lt;a href=evil&gt;04/15/2026&lt;/a&gt;" in html


def test_html_single_state_and_true_multi_state_keep_legacy_labels():
    single = _base_data()
    single["tax_summary"]["state_sentences"] = [
        {
            "state_name": "California",
            "state_abbreviation": "CA",
            "display_label": None,
            "sentence": "**Refund** of **$297**.",
        }
    ]
    multi = _base_data()
    multi["tax_summary"]["state_sentences"] = [
        {
            "state_name": "California",
            "state_abbreviation": "CA",
            "display_label": None,
            "sentence": "**Balance due** of **$29**.",
        },
        {
            "state_name": "North Carolina",
            "state_abbreviation": "NC",
            "display_label": None,
            "sentence": "**Refund** of **$95**.",
        },
    ]

    single_html = generate_email_html(single)
    multi_html = generate_email_html(multi)

    assert "State Income Tax:" in single_html
    assert "CA State Income Tax:" not in single_html
    assert "CA State Income Tax:" in multi_html
    assert "NC State Income Tax:" in multi_html


def _accupuncture_data():
    data = _base_data()
    data["return_type"] = "S-Corporation (1120S)"
    data["tax_summary"] = {
        "federal_sentence": "**No tax** is payable with the filing of this return.",
        "state_sentences": [
            {
                "state_name": "California",
                "state_abbreviation": "CA",
                "display_label": None,
                "sentence": "**No tax** is payable with the filing of this return.",
            }
        ],
    }
    data["estimated_payments"] = [
        {"date": "04/15/2026", "federal": 0, "state": 2800, "state_name": "California"},
        {"date": "06/15/2026", "federal": 0, "state": 3800, "state_name": "California"},
        {"date": "12/15/2026", "federal": 0, "state": 2800, "state_name": "California"},
    ]
    data["pte_payments"] = [
        {
            "sentence": (
                "The 2026 Pass-Through Entity Elective Tax Estimate **balance due** "
                "of **$28,805** will be withdrawn on **June 15, 2026**."
            )
        }
    ]
    return data


def test_html_accupuncture_estimates_and_pte_remain_unchanged():
    data = _accupuncture_data()

    html = generate_email_html(data)

    assert "State Income Tax:" in html
    assert "$2,800" in html
    assert "$3,800" in html
    assert "$28,805" in html


def _docx_text(path):
    doc = Document(path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _docx_table_text(path):
    doc = Document(path)
    return "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )


def test_docx_combines_cny_brand_and_correct_im_labels(tmp_path):
    output = tmp_path / "im.docx"

    generate_email_docx(_im_data(), output)
    text = _docx_text(output)

    federal_line = (
        "Federal Income Tax: No tax is payable with the filing of this return. "
        "Refund of $8,647 will be deposited to your account."
    )
    state_line = (
        "State Income Tax: No tax is payable with the filing of this return. "
        "Refund of $7,447 will be deposited to your account."
    )
    metro_line = (
        "OR Metro Income Tax: No tax is payable with the filing of this return. "
        "Refund of $349 will be deposited to your account."
    )

    assert "Your 2025 Income Tax Return and CNY invoice are now available" in text
    assert federal_line in text
    assert state_line in text
    assert metro_line in text
    assert text.index(federal_line) < text.index(state_line) < text.index(metro_line)
    assert "OR State Income Tax" not in text
    assert 'Enter "cnyllp" as the subdomain' in text


def test_hong_refunds_remain_associated_with_the_correct_labels(tmp_path):
    output = tmp_path / "hong.docx"
    data = _hong_data()

    html = generate_email_html(data)
    generate_email_docx(data, output)
    docx_text = _docx_text(output)

    assert (
        "Federal Income Tax: No tax is payable with the filing of this return. "
        "<strong>Refund</strong> of <strong>$951</strong> will be deposited to your account."
    ) in html
    assert (
        "State Income Tax: No tax is payable with the filing of this return. "
        "<strong>Refund</strong> of <strong>$297</strong> will be deposited to your account."
    ) in html
    assert (
        "Federal Income Tax: No tax is payable with the filing of this return. "
        "Refund of $951 will be deposited to your account."
    ) in docx_text
    assert (
        "State Income Tax: No tax is payable with the filing of this return. "
        "Refund of $297 will be deposited to your account."
    ) in docx_text


def test_docx_treats_display_label_as_plain_text_and_markdown_only_in_sentence(tmp_path):
    output = tmp_path / "literal-label.docx"
    data = _base_data()
    data["tax_summary"]["state_sentences"] = [
        {
            "state_name": "Oregon Metro",
            "state_abbreviation": "OR",
            "display_label": "OR **Metro** Income Tax",
            "sentence": "**Refund** of **$1**.",
        }
    ]

    generate_email_docx(data, output)
    doc = Document(output)
    paragraph = next(p for p in doc.paragraphs if "Metro" in p.text)

    assert paragraph.text == "OR **Metro** Income Tax: Refund of $1."
    assert [(run.text, run.bold) for run in paragraph.runs] == [
        ("OR **Metro** Income Tax: ", False),
        ("Refund", True),
        (" of ", False),
        ("$1", True),
        (".", False),
    ]


def test_docx_accupuncture_estimates_and_pte_remain_unchanged(tmp_path):
    output = tmp_path / "accupuncture.docx"

    generate_email_docx(_accupuncture_data(), output)
    text = _docx_text(output)
    table_text = _docx_table_text(output)

    assert "State Income Tax:" in text
    assert "$28,805" in text
    assert "$2,800" in table_text
    assert "$3,800" in table_text
    assert "04/15/2026" in table_text
    assert "06/15/2026" in table_text
    assert "12/15/2026" in table_text


@pytest.mark.parametrize("firm", ["CNY LLP", "ABC LLP", None, ""])
def test_docx_invoice_brand_is_global_and_does_not_mutate_input(firm, tmp_path):
    output = tmp_path / "email.docx"
    data = _base_data(firm)
    before = copy.deepcopy(data)

    generate_email_docx(data, output)
    text = _docx_text(output)

    assert "Your 2025 Income Tax Return and CNY invoice are now available" in text
    assert "CNY LLP invoice" not in text
    assert "ABC LLP invoice" not in text
    assert "incoive" not in text
    assert data == before


@pytest.mark.parametrize("firm_value", ["missing", None])
def test_docx_handles_missing_or_null_firm_object(firm_value, tmp_path):
    output = tmp_path / "email.docx"
    data = _base_data()
    if firm_value == "missing":
        data.pop("cpa_firm")
    else:
        data["cpa_firm"] = None

    generate_email_docx(data, output)
    text = _docx_text(output)

    assert "CNY invoice" in text
    assert "Enter the firm subdomain" in text


def test_docx_handles_firm_object_without_name(tmp_path):
    output = tmp_path / "email.docx"
    data = _base_data()
    data["cpa_firm"].pop("name")

    generate_email_docx(data, output)
    text = _docx_text(output)

    assert "CNY invoice" in text
    assert 'Enter "cnyllp" as the subdomain' in text


def test_docx_single_state_and_true_multi_state_keep_legacy_labels(tmp_path):
    single = _base_data()
    single["tax_summary"]["state_sentences"] = [
        {
            "state_name": "California",
            "state_abbreviation": "CA",
            "display_label": None,
            "sentence": "**Refund** of **$297**.",
        }
    ]
    multi = _base_data()
    multi["tax_summary"]["state_sentences"] = [
        {
            "state_name": "California",
            "state_abbreviation": "CA",
            "display_label": None,
            "sentence": "**Balance due** of **$29**.",
        },
        {
            "state_name": "North Carolina",
            "state_abbreviation": "NC",
            "display_label": None,
            "sentence": "**Refund** of **$95**.",
        },
    ]
    single_output = tmp_path / "single.docx"
    multi_output = tmp_path / "multi.docx"

    generate_email_docx(single, single_output)
    generate_email_docx(multi, multi_output)
    single_text = _docx_text(single_output)
    multi_text = _docx_text(multi_output)

    assert "State Income Tax:" in single_text
    assert "CA State Income Tax:" not in single_text
    assert "CA State Income Tax:" in multi_text
    assert "NC State Income Tax:" in multi_text
